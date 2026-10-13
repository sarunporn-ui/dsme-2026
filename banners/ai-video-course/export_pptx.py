"""Export แบนเนอร์ใน index.html เป็น PowerPoint ที่แก้ไขได้

- ข้อความทุกจุด -> Text box จริง (ฟอนต์ LINE Seed Sans TH) วางตำแหน่ง/สี/ขนาดตาม HTML
- รูป/การ์ด/วงกลม/ตั๋ว/ปุ่ม -> รูป PNG โปร่งใสแยกชิ้น (ย้าย/ลบ/เปลี่ยนได้)
- 1 ไฟล์ต่อ 1 ขนาด (1:1, 4:5, 16:9) ไฟล์ละ 3 สไลด์ (A/B/C)

ใช้: python3 export_pptx.py   (ต้องมี playwright + python-pptx)
"""
import io
import math
import os
import re

from lxml import etree
from PIL import Image
from playwright.sync_api import sync_playwright
from pptx import Presentation
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, "pptx")
CHROME = os.environ.get("CHROME", "/opt/pw-browsers/chromium")
PX = 9525  # EMU ต่อ 1px (96 dpi)

SIZES = {
    "1x1": (1080, 1080, ["a11", "b11", "c11"]),
    "4x5": (1080, 1350, ["a45", "b45", "c45"]),
    "16x9": (1440, 810, ["a169", "b169", "c169"]),
}
NAMES = {"a": "A · Classic", "b": "B · Spotlight", "c": "C · Ticket"}

FONT = "LINE Seed Sans TH"
FONT_HEAVY = "LINE Seed Sans TH Heavy"

# ชิ้นที่ส่งออกเป็นรูป (เอาชั้นนอกสุด) — ข้อความข้างในถูกซ่อนแล้วไปเป็น text box แทน
VISUAL = ".photo, img, .tool, .stat, .price, .badge, .bar2, .ticket, .frame, .btn"

EXTRACT_JS = r"""
() => {
  const banner = document.querySelector('.banner');
  const B = banner.getBoundingClientRect();
  const EXCL = '.tool, .logo';
  const cs = el => getComputedStyle(el);
  const isInline = n => n.nodeType === 1 && cs(n).display === 'inline' &&
                        !['IMG', 'svg', 'SVG'].includes(n.tagName);

  // ---------- ข้อความ ----------
  function collect(node, runs, textNodes) {
    for (const ch of node.childNodes) {
      if (ch.nodeType === 3) {
        const t = ch.textContent.replace(/\s+/g, ' ');
        if (!t.trim() && !t.includes(' ')) continue;
        const s = cs(ch.parentElement);
        runs.push({t, color: s.color, size: parseFloat(s.fontSize), weight: parseInt(s.fontWeight),
                   strike: s.textDecorationLine.includes('line-through'),
                   ls: parseFloat(s.letterSpacing) || 0, shadow: s.textShadow !== 'none'});
        if (t.trim()) textNodes.push(ch);
      } else if (ch.nodeName === 'BR') {
        runs.push({br: true});
      } else if (isInline(ch)) {
        collect(ch, runs, textNodes);
      }
    }
  }
  function bbox(nodes) {
    let x0 = 1e9, y0 = 1e9, x1 = -1e9, y1 = -1e9;
    for (const n of nodes) {
      const r = document.createRange(); r.selectNodeContents(n);
      for (const q of r.getClientRects()) {
        x0 = Math.min(x0, q.left); y0 = Math.min(y0, q.top);
        x1 = Math.max(x1, q.right); y1 = Math.max(y1, q.bottom);
      }
    }
    return {x: x0 - B.left, y: y0 - B.top, w: x1 - x0, h: y1 - y0};
  }
  function rotAncestor(el) {
    for (let a = el; a && a !== banner; a = a.parentElement) {
      const t = cs(a).transform;
      if (t && t !== 'none') {
        const m = t.match(/matrix\(([^)]+)\)/);
        if (!m) continue;
        const [ma, mb] = m[1].split(',').map(parseFloat);
        const deg = Math.atan2(mb, ma) * 180 / Math.PI;
        if (Math.abs(deg) > 0.1) return {el: a, deg};
      }
    }
    return null;
  }
  // แยกเป็น 1 กล่องต่อ 1 บรรทัด วางด้วย baseline (ฟอนต์นี้: ascent 1.144em, สูงรวม 1.584em)
  const ASC = 1.144, LINE = 1.584;
  const texts = [];
  for (const el of banner.querySelectorAll('*')) {
    if (el.closest(EXCL) || isInline(el) || ['IMG', 'svg', 'SVG'].includes(el.tagName)) continue;
    if (el.closest('svg')) continue;
    const runs = [], nodes = [];
    collect(el, runs, nodes);
    if (!nodes.length) continue;
    // จับคู่ run กับ text node ตามลำดับ แล้วตัดเป็นบรรทัดที่ <br>
    const lines = [[]];
    let ni = 0;
    for (const r of runs) {
      if (r.br) { lines.push([]); continue; }
      if (r.t.trim()) { r.node = nodes[ni++]; }
      lines[lines.length - 1].push(r);
    }
    const s = cs(el);
    const rot = rotAncestor(el);
    const old = rot ? rot.el.style.transform : null;
    if (rot) rot.el.style.transform = 'none';
    let ac = null;
    if (rot) { const ar = rot.el.getBoundingClientRect();
               ac = {x: ar.left - B.left + ar.width / 2, y: ar.top - B.top + ar.height / 2}; }
    for (const ln of lines) {
      const ns = ln.filter(r => r.node);
      if (!ns.length) continue;
      const bb = bbox(ns.map(r => r.node));
      const big = ns.reduce((a, b) => (b.size > a.size ? b : a));
      const rr = document.createRange(); rr.selectNodeContents(big.node);
      const baseline = rr.getClientRects()[0].top - B.top + ASC * big.size;
      const maxfs = Math.max(...ln.map(r => r.size || 0));
      let box = {x: bb.x, y: baseline - ASC * maxfs, w: bb.w, h: LINE * maxfs, rot: 0};
      if (rot) {
        const tc = {x: box.x + box.w / 2, y: box.y + box.h / 2};
        const a = rot.deg * Math.PI / 180, dx = tc.x - ac.x, dy = tc.y - ac.y;
        const cx = ac.x + dx * Math.cos(a) - dy * Math.sin(a);
        const cy = ac.y + dx * Math.sin(a) + dy * Math.cos(a);
        box = {x: cx - box.w / 2, y: cy - box.h / 2, w: box.w, h: box.h, rot: rot.deg};
      }
      ln.forEach(r => delete r.node);
      texts.push({...box, runs: ln, align: s.textAlign});
    }
    if (rot) rot.el.style.transform = old;
  }

  // ---------- ชิ้นภาพ ----------
  const VIS = '__VISUAL__';
  const all = [...banner.querySelectorAll(VIS)];
  const visuals = all.filter(el => !all.some(o => o !== el && o.contains(el)));
  visuals.forEach((el, i) => el.dataset.xp = i);
  return {texts, nvis: visuals.length};
}
""".replace("__VISUAL__", VISUAL)

HIDE_ALL = """
.banner *{visibility:hidden!important}
"""
ONLY_TARGET = """
body,.banner{background:transparent!important}
.banner *{visibility:hidden!important}
.xp-t,.xp-t *{visibility:visible!important}
.xp-t:not(.tool),.xp-t:not(.tool) *{color:transparent!important;text-shadow:none!important}
.banner .xp-t svg,.banner .xp-t svg *{color:#fff!important}
"""


def rgb(css):
    m = re.match(r"rgba?\(([^)]+)\)", css)
    v = [float(x) for x in m.group(1).split(",")]
    return "%02X%02X%02X" % tuple(int(round(c)) for c in v[:3])


def png_layer(page, w, h):
    png = page.screenshot(clip={"x": 0, "y": 0, "width": w, "height": h}, omit_background=True)
    im = Image.open(io.BytesIO(png)).convert("RGBA")
    box = im.getchannel("A").getbbox()
    if not box:
        return None, None
    return im.crop(box), box


def add_picture(slide, im, x, y, scale):
    buf = io.BytesIO()
    im.save(buf, "PNG")
    buf.seek(0)
    slide.shapes.add_picture(buf, Emu(int(x / scale * PX)), Emu(int(y / scale * PX)),
                             Emu(int(im.width / scale * PX)), Emu(int(im.height / scale * PX)))


def set_run(run, r):
    run.text = r["t"]
    rPr = run._r.get_or_add_rPr()
    rPr.set("sz", str(int(round(r["size"] * 0.75 * 100))))
    heavy = r["weight"] >= 800
    rPr.set("b", "1" if 600 <= r["weight"] < 800 else "0")
    if r["strike"]:
        rPr.set("strike", "sngStrike")
    if abs(r["ls"]) > 0.01:
        rPr.set("spc", str(int(round(r["ls"] * 0.75 * 100))))
    fill = etree.SubElement(rPr, qn("a:solidFill"))
    etree.SubElement(fill, qn("a:srgbClr"), val=rgb(r["color"]))
    if r["shadow"]:
        eff = etree.SubElement(rPr, qn("a:effectLst"))
        sh = etree.SubElement(eff, qn("a:outerShdw"), blurRad="76200", dist="38100", dir="5400000",
                              algn="t", rotWithShape="0")
        c = etree.SubElement(sh, qn("a:srgbClr"), val="000000")
        etree.SubElement(c, qn("a:alpha"), val="45000")
    face = FONT_HEAVY if heavy else FONT
    for tag in ("a:latin", "a:ea", "a:cs"):
        etree.SubElement(rPr, qn(tag), typeface=face)


def add_text(slide, t):
    pad = 6  # px กันข้อความตกขอบ (ฟอนต์ใน PowerPoint กว้างกว่าเบราว์เซอร์เล็กน้อย)
    x, y, w, h = t["x"] - pad, t["y"], t["w"] + 2 * pad, t["h"]
    tb = slide.shapes.add_textbox(Emu(int(x * PX)), Emu(int(y * PX)), Emu(int(w * PX)), Emu(int(h * PX)))
    if t["rot"]:
        tb.rotation = t["rot"]
    tf = tb.text_frame
    tf.word_wrap = False
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.vertical_anchor = MSO_ANCHOR.TOP
    tf.margin_left = tf.margin_right = Emu(pad * PX)
    tf.margin_top = tf.margin_bottom = 0
    align = {"center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT, "end": PP_ALIGN.RIGHT}.get(t["align"], PP_ALIGN.LEFT)

    ln = [dict(r) for r in t["runs"]]
    ln[0]["t"] = ln[0]["t"].lstrip()  # ตัดช่องว่างหัว/ท้ายบรรทัดแบบเบราว์เซอร์
    ln[-1]["t"] = ln[-1]["t"].rstrip()
    p = tf.paragraphs[0]
    p.alignment = align
    p.line_spacing = 1.0
    for r in ln:
        if r["t"]:
            set_run(p.add_run(), r)


def build():
    os.makedirs(OUT, exist_ok=True)
    url = "file://" + os.path.join(HERE, "index.html")
    with sync_playwright() as p:
        browser = p.chromium.launch(executable_path=CHROME, args=["--allow-file-access-from-files"])
        for key, (w, h, ids) in SIZES.items():
            prs = Presentation()
            prs.slide_width, prs.slide_height = Emu(w * PX), Emu(h * PX)
            for bid in ids:
                page = browser.new_page(viewport={"width": w, "height": h}, device_scale_factor=2)
                page.goto(f"{url}#{bid}")
                page.evaluate("document.fonts.ready")
                page.wait_for_timeout(300)
                data = page.evaluate(EXTRACT_JS)
                slide = prs.slides.add_slide(prs.slide_layouts[6])
                slide.name = NAMES[bid[0]]

                # 1) พื้นหลัง
                st = page.add_style_tag(content=HIDE_ALL)
                bg = Image.open(io.BytesIO(page.screenshot(clip={"x": 0, "y": 0, "width": w, "height": h}))).convert("RGB")
                st.evaluate("e => e.remove()")
                add_picture(slide, bg, 0, 0, 2)

                # 2) ชิ้นภาพทีละชิ้น ตามลำดับซ้อนใน HTML
                st = page.add_style_tag(content=ONLY_TARGET)
                for i in range(data["nvis"]):
                    page.evaluate(f"document.querySelector('[data-xp=\"{i}\"]').classList.add('xp-t')")
                    im, box = png_layer(page, w, h)
                    page.evaluate(f"document.querySelector('[data-xp=\"{i}\"]').classList.remove('xp-t')")
                    if im:
                        add_picture(slide, im, box[0], box[1], 2)
                st.evaluate("e => e.remove()")

                # 3) ข้อความ (แก้ไขได้)
                for t in data["texts"]:
                    add_text(slide, t)
                page.close()
            path = os.path.join(OUT, f"1Day-AI-ตัดคลิป_{key}.pptx")
            prs.save(path)
            print("saved", path)
        browser.close()


if __name__ == "__main__":
    build()
